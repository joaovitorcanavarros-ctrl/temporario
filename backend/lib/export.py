"""Geração de arquivos CSV (UTF-8 com BOM, separador ';' para Excel pt-BR) e XLSX."""

import csv
import io
from typing import Any
from urllib.parse import quote

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _fmt(v: Any) -> Any:
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, str) and len(v) == 10 and v[4] == "-" and v[7] == "-":
        return f"{v[8:10]}/{v[5:7]}/{v[0:4]}"
    return "" if v is None else v


def file_response(titulo: str, colunas: list[tuple[str, str]], linhas: list[dict], formato: str, nome: str) -> StreamingResponse:
    if formato == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = titulo[:31]
        ws.append([label for _, label in colunas])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E3A8A")
        for li in linhas:
            ws.append([_fmt(li.get(k)) for k, _ in colunas])
        for i, (_, label) in enumerate(colunas, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(60, len(label) + 6))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"{nome}.xlsx"
        body: Any = buf
    else:
        sio = io.StringIO()
        w = csv.writer(sio, delimiter=";")
        w.writerow([label for _, label in colunas])
        for li in linhas:
            w.writerow([_fmt(li.get(k)) for k, _ in colunas])
        body = io.BytesIO(("\ufeff" + sio.getvalue()).encode("utf-8"))
        media = "text/csv; charset=utf-8"
        filename = f"{nome}.csv"
    return StreamingResponse(
        body, media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
